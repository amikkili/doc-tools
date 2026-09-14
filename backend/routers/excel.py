import io
from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from .utils import check_size, safe_error, validate_content_type

router = APIRouter()


@router.post("/to-pdf")
async def excel_to_pdf(file: UploadFile = File(...)):
    """Convert XLSX/XLS to PDF — renders all sheets as tables."""
    try:
        data = await file.read()
        check_size(data)
        validate_content_type(data, "xlsx", file.filename or "File")
        pdf_bytes = _excel_to_pdf(data, file.filename or "file.xlsx")
        return StreamingResponse(
            io.BytesIO(pdf_bytes),
            media_type="application/pdf",
            headers={"Content-Disposition": 'attachment; filename="converted.pdf"'},
        )
    except Exception as e:
        raise HTTPException(500, safe_error(e))


def _excel_to_pdf(data: bytes, filename: str) -> bytes:
    import shutil, subprocess, tempfile, os

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        ext = ".xlsx" if filename.endswith(".xlsx") else ".xls"
        with tempfile.TemporaryDirectory() as tmp_dir:
            inp = os.path.join(tmp_dir, f"input{ext}")
            with open(inp, "wb") as f:
                f.write(data)
            result = subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp_dir, inp],
                capture_output=True, timeout=60,
            )
            if result.returncode != 0:
                raise RuntimeError(result.stderr.decode())
            pdf_path = inp.rsplit(".", 1)[0] + ".pdf"
            with open(pdf_path, "rb") as f:
                return f.read()

    return _reportlab_excel(data)


def _reportlab_excel(data: bytes) -> bytes:
    """Fallback: render spreadsheet as a table using reportlab."""
    import openpyxl
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    styles = getSampleStyleSheet()
    buf = io.BytesIO()
    pdf = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=10*mm, rightMargin=10*mm,
                            topMargin=10*mm, bottomMargin=10*mm)
    story = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        story.append(Paragraph(sheet_name, styles["Heading2"]))
        story.append(Spacer(1, 4))

        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append([str(c) if c is not None else "" for c in row])

        if rows:
            col_count = max(len(r) for r in rows)
            for r in rows:
                while len(r) < col_count:
                    r.append("")
            available = landscape(A4)[0] - 20*mm
            col_w = available / max(col_count, 1)
            t = Table(rows, colWidths=[col_w] * col_count)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#e74c3c")),
                ("TEXTCOLOR", (0,0), (-1,0), colors.white),
                ("FONTSIZE", (0,0), (-1,-1), 7),
                ("GRID", (0,0), (-1,-1), 0.5, colors.grey),
                ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#fef2f2")]),
                ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
                ("WORDWRAP", (0,0), (-1,-1), True),
            ]))
            story.append(t)
        story.append(Spacer(1, 10))

    pdf.build(story)
    return buf.getvalue()
