import io
import zipfile
from pathlib import Path
from typing import List, Optional

from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import Response, StreamingResponse
from .utils import check_size, safe_error, validate_content_type

router = APIRouter()


# ── helpers ────────────────────────────────────────────────────────────────

def stream(buf: io.BytesIO, filename: str, media_type: str = "application/pdf") -> StreamingResponse:
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ── Merge ──────────────────────────────────────────────────────────────────

@router.post("/merge")
async def merge_pdf(files: List[UploadFile] = File(...)):
    if len(files) < 2:
        raise HTTPException(400, "At least 2 PDF files required")
    try:
        from pypdf import PdfWriter
        writer = PdfWriter()
        for f in files:
            data = await f.read()
            check_size(data, f.filename or "File")
            validate_content_type(data, "pdf", f.filename or "File")
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            for page in reader.pages:
                writer.add_page(page)
        buf = io.BytesIO()
        writer.write(buf)
        return stream(buf, "merged.pdf")
    except Exception as e:
        raise HTTPException(500, safe_error(e))


# ── Split ──────────────────────────────────────────────────────────────────

def _parse_ranges(pages_str: str, total: int) -> List[List[int]]:
    """Parse '1-3,5,7-9' into list of page-index lists (0-based)."""
    parts = [p.strip() for p in pages_str.split(",") if p.strip()]
    groups = []
    for part in parts:
        if "-" in part:
            a, b = part.split("-", 1)
            groups.append(list(range(int(a) - 1, min(int(b), total))))
        else:
            groups.append([int(part) - 1])
    return groups


@router.post("/split")
async def split_pdf(file: UploadFile = File(...), pages: Optional[str] = Form(None)):
    try:
        from pypdf import PdfReader, PdfWriter
        data = await file.read()
        check_size(data)
        validate_content_type(data, "pdf")
        reader = PdfReader(io.BytesIO(data))
        total = len(reader.pages)

        if pages and pages.strip():
            groups = _parse_ranges(pages, total)
        else:
            groups = [[i] for i in range(total)]

        zip_buf = io.BytesIO()
        with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
            for idx, page_ids in enumerate(groups, 1):
                writer = PdfWriter()
                for pid in page_ids:
                    if 0 <= pid < total:
                        writer.add_page(reader.pages[pid])
                part_buf = io.BytesIO()
                writer.write(part_buf)
                zf.writestr(f"part_{idx}.pdf", part_buf.getvalue())

        return stream(zip_buf, "split.zip", "application/zip")
    except Exception as e:
        raise HTTPException(500, safe_error(e))


# ── Compress ───────────────────────────────────────────────────────────────

@router.post("/compress")
async def compress_pdf(file: UploadFile = File(...), quality: str = Form("medium")):
    try:
        import fitz  # PyMuPDF

        dpi_map = {"low": 72, "medium": 120, "high": 200}
        dpi = dpi_map.get(quality, 120)

        data = await file.read()
        check_size(data)
        doc = fitz.open(stream=data, filetype="pdf")

        out_doc = fitz.open()
        for page in doc:
            pix = page.get_pixmap(dpi=dpi)
            img_pdf = fitz.open()
            img_page = img_pdf.new_page(width=page.rect.width, height=page.rect.height)
            img_page.insert_image(img_page.rect, pixmap=pix)
            out_doc.insert_pdf(img_pdf)

        buf = io.BytesIO(out_doc.tobytes(deflate=True))
        return stream(buf, "compressed.pdf")
    except Exception as e:
        raise HTTPException(500, safe_error(e))


# ── Rotate ─────────────────────────────────────────────────────────────────

@router.post("/rotate")
async def rotate_pdf(file: UploadFile = File(...), angle: int = Form(90)):
    if angle not in (90, 180, 270):
        raise HTTPException(400, "Angle must be 90, 180 or 270")
    try:
        from pypdf import PdfReader, PdfWriter
        data = await file.read()
        check_size(data)
        reader = PdfReader(io.BytesIO(data))
        writer = PdfWriter()
        for page in reader.pages:
            page.rotate(angle)
            writer.add_page(page)
        buf = io.BytesIO()
        writer.write(buf)
        return stream(buf, "rotated.pdf")
    except Exception as e:
        raise HTTPException(500, safe_error(e))


# ── Watermark ──────────────────────────────────────────────────────────────

@router.post("/watermark")
async def watermark_pdf(
    file: UploadFile = File(...),
    text: str = Form("CONFIDENTIAL"),
    opacity: float = Form(0.3),
    color: str = Form("#FF0000"),
):
    try:
        import fitz

        def hex_to_rgb(h: str):
            h = h.lstrip("#")
            return tuple(int(h[i:i+2], 16) / 255 for i in (0, 2, 4))

        r, g, b = hex_to_rgb(color)
        data = await file.read()
        check_size(data)
        doc = fitz.open(stream=data, filetype="pdf")

        for page in doc:
            w, h = page.rect.width, page.rect.height
            page.insert_text(
                (w * 0.1, h * 0.55),
                text,
                fontsize=min(w, h) * 0.1,
                color=(r, g, b),
                fill_opacity=opacity,
                morph=(fitz.Point(w / 2, h / 2), fitz.Matrix(45)),
                overlay=True,
            )

        buf = io.BytesIO(doc.tobytes())
        return stream(buf, "watermarked.pdf")
    except Exception as e:
        raise HTTPException(500, safe_error(e))


# ── Protect ────────────────────────────────────────────────────────────────

@router.post("/protect")
async def protect_pdf(file: UploadFile = File(...), password: str = Form(...)):
    try:
        from pypdf import PdfReader, PdfWriter
        data = await file.read()
        check_size(data)
        reader = PdfReader(io.BytesIO(data))
        writer = PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        writer.encrypt(password)
        buf = io.BytesIO()
        writer.write(buf)
        return stream(buf, "protected.pdf")
    except Exception as e:
        raise HTTPException(500, safe_error(e))


# ── PDF → Word ─────────────────────────────────────────────────────────────

@router.post("/to-word")
async def pdf_to_word(file: UploadFile = File(...)):
    try:
        from pdf2docx import Converter
        import tempfile, os

        data = await file.read()
        check_size(data)
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp_pdf:
            tmp_pdf.write(data)
            tmp_pdf_path = tmp_pdf.name

        tmp_docx = tmp_pdf_path.replace(".pdf", ".docx")
        try:
            cv = Converter(tmp_pdf_path)
            cv.convert(tmp_docx, start=0, end=None)
            cv.close()
            with open(tmp_docx, "rb") as f:
                docx_bytes = f.read()
        finally:
            os.unlink(tmp_pdf_path)
            if Path(tmp_docx).exists():
                os.unlink(tmp_docx)

        buf = io.BytesIO(docx_bytes)
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": 'attachment; filename="converted.docx"'},
        )
    except Exception as e:
        raise HTTPException(500, safe_error(e))


# ── PDF → JPG ──────────────────────────────────────────────────────────────

@router.post("/to-jpg")
async def pdf_to_jpg(file: UploadFile = File(...), dpi: int = Form(150)):
    try:
        import fitz

        data = await file.read()
        check_size(data)
        doc = fitz.open(stream=data, filetype="pdf")
        mat = fitz.Matrix(dpi / 72, dpi / 72)

        zip_buf = io.BytesIO()
        with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
            for i, page in enumerate(doc):
                pix = page.get_pixmap(matrix=mat)
                img_bytes = pix.tobytes("jpeg")
                zf.writestr(f"page_{i+1:03d}.jpg", img_bytes)

        return stream(zip_buf, "pdf-images.zip", "application/zip")
    except Exception as e:
        raise HTTPException(500, safe_error(e))


# ── Images → PDF ───────────────────────────────────────────────────────────

@router.post("/from-images")
async def images_to_pdf(
    files: List[UploadFile] = File(...),
    orientation: str = Form("portrait"),
):
    try:
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.platypus import SimpleDocTemplate, Image as RLImage
        from reportlab.lib.units import mm

        page_size = landscape(A4) if orientation == "landscape" else A4
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=page_size,
                                leftMargin=10*mm, rightMargin=10*mm,
                                topMargin=10*mm, bottomMargin=10*mm)
        story = []
        pw = page_size[0] - 20*mm
        ph = page_size[1] - 20*mm

        from PIL import Image as PILImage
        for f in files:
            img_data = await f.read()
            img = PILImage.open(io.BytesIO(img_data))
            iw, ih = img.size
            ratio = min(pw / iw, ph / ih)
            rli = RLImage(io.BytesIO(img_data), width=iw*ratio, height=ih*ratio)
            story.append(rli)

        doc.build(story)
        return stream(buf, "images.pdf")
    except Exception as e:
        raise HTTPException(500, safe_error(e))


# ── PowerPoint → PDF ───────────────────────────────────────────────────────

@router.post("/from-pptx")
async def pptx_to_pdf(file: UploadFile = File(...)):
    try:
        import shutil, subprocess, tempfile, os

        data = await file.read()
        check_size(data)

        soffice = shutil.which('soffice') or shutil.which('libreoffice')
        if not soffice:
            raise HTTPException(
                503,
                "PPTX→PDF requires LibreOffice on the server. "
                "Install with: sudo apt-get install -y libreoffice"
            )

        with tempfile.TemporaryDirectory() as tmpdir:
            src = os.path.join(tmpdir, file.filename or "input.pptx")
            with open(src, "wb") as f:
                f.write(data)

            # Each conversion gets its own isolated profile dir to prevent
            # LibreOffice lock-file conflicts that silently produce blank PDFs.
            profile_dir = os.path.join(tmpdir, "lo_profile")
            os.makedirs(profile_dir, exist_ok=True)
            user_install = f"file://{profile_dir}"

            proc = subprocess.run(
                [
                    soffice,
                    f"-env:UserInstallation={user_install}",
                    "--headless",
                    "--norestore",
                    "--convert-to", "pdf",
                    "--outdir", tmpdir,
                    src,
                ],
                capture_output=True, text=True, timeout=120,
                env={**os.environ, "HOME": tmpdir},
            )
            if proc.returncode != 0:
                raise HTTPException(500, f"Conversion failed: {proc.stderr[:500]}")

            pdf_path = os.path.join(tmpdir, Path(src).stem + ".pdf")
            if not os.path.exists(pdf_path):
                raise HTTPException(500, "LibreOffice produced no output file")

            pdf_bytes = open(pdf_path, "rb").read()
            if len(pdf_bytes) < 500:
                raise HTTPException(500, f"Conversion produced an empty PDF (stderr: {proc.stderr[:300]})")

        stem = Path(file.filename or "presentation").stem
        return StreamingResponse(
            io.BytesIO(pdf_bytes),
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="{stem}.pdf"'},
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(500, safe_error(e))


# ── PDF → PowerPoint ───────────────────────────────────────────────────────

@router.post("/to-pptx")
async def pdf_to_pptx(file: UploadFile = File(...)):
    try:
        import fitz
        from pptx import Presentation
        from pptx.util import Emu

        data = await file.read()
        check_size(data)
        doc = fitz.open(stream=data, filetype="pdf")

        if len(doc) == 0:
            raise HTTPException(400, "PDF has no pages")

        prs = Presentation()

        # Match slide size to PDF page (1 pt = 12700 EMU)
        first = doc[0]
        slide_w = Emu(int(first.rect.width  * 12700))
        slide_h = Emu(int(first.rect.height * 12700))
        prs.slide_width  = slide_w
        prs.slide_height = slide_h

        # Find a truly blank layout (no placeholders)
        blank_layout = next(
            (l for l in prs.slide_layouts if len(l.placeholders) == 0),
            prs.slide_layouts[6],
        )

        for page in doc:
            # Render at 150 DPI and keep image bytes in memory — no temp files
            pix = page.get_pixmap(matrix=fitz.Matrix(150 / 72, 150 / 72))
            img_buf = io.BytesIO(pix.tobytes("png"))

            slide = prs.slides.add_slide(blank_layout)
            slide.shapes.add_picture(
                img_buf,
                left=0, top=0,
                width=slide_w,
                height=slide_h,
            )

        stem = Path(file.filename or "presentation").stem
        buf = io.BytesIO()
        prs.save(buf)
        return Response(
            content=buf.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            headers={"Content-Disposition": f'attachment; filename="{stem}.pptx"'},
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(500, safe_error(e))


# ── PDF → Excel ────────────────────────────────────────────────────────────

@router.post("/to-excel")
async def pdf_to_excel(file: UploadFile = File(...)):
    try:
        import tempfile, os
        import openpyxl

        data = await file.read()
        check_size(data)

        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
            tmp.write(data)
            tmp_path = tmp.name

        try:
            # Try tabula for structured table extraction
            try:
                import tabula
                dfs = tabula.read_pdf(tmp_path, pages="all", multiple_tables=True, silent=True)
                dfs = [df for df in dfs if not df.empty and df.shape[1] > 1 and df.shape[0] > 0]
            except Exception:
                dfs = []

            wb = openpyxl.Workbook()
            wb.remove(wb.active)

            if dfs:
                for i, df in enumerate(dfs):
                    ws = wb.create_sheet(title=f"Table {i + 1}")
                    _pdf_excel_styled_table(ws, df)
            else:
                # No tables — extract text line-by-line with PyMuPDF
                ws = wb.create_sheet(title="Content")
                _pdf_excel_text_layout(ws, tmp_path)

        finally:
            os.unlink(tmp_path)

        fname = (file.filename or "extracted").rsplit(".", 1)[0]
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{fname}.xlsx"'},
        )
    except Exception as e:
        raise HTTPException(500, safe_error(e))


def _xlsx_header_style(cell, bg="1E3A5F"):
    from openpyxl.styles import Font, PatternFill, Alignment
    cell.font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    cell.fill = PatternFill("solid", fgColor=bg)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _xlsx_data_style(cell, even=False):
    from openpyxl.styles import Font, PatternFill, Alignment
    cell.font = Font(name="Calibri", size=10)
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    if even:
        cell.fill = PatternFill("solid", fgColor="EFF4FB")


def _xlsx_thin_border(ws):
    from openpyxl.styles import Border, Side
    thin = Side(style="thin", color="D0D7DE")
    bdr = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows():
        for cell in row:
            cell.border = bdr


def _xlsx_auto_widths(ws, min_w=10, max_w=55):
    from openpyxl.utils import get_column_letter
    for col in ws.columns:
        best = 0
        letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                best = max(best, len(str(cell.value or "")))
            except Exception:
                pass
        ws.column_dimensions[letter].width = min(max(best + 2, min_w), max_w)


def _pdf_excel_styled_table(ws, df):
    """Write a pandas DataFrame as a beautifully formatted Excel sheet."""
    import math

    cols = list(df.columns)
    n_cols = len(cols)

    # Header
    ws.append(cols)
    for c in range(1, n_cols + 1):
        _xlsx_header_style(ws.cell(row=1, column=c))
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"

    # Data rows — clean NaN
    for row_idx, row in enumerate(df.itertuples(index=False), start=2):
        values = []
        for v in row:
            try:
                if v is None or (isinstance(v, float) and math.isnan(v)):
                    values.append("")
                else:
                    values.append(str(v))
            except Exception:
                values.append("")
        ws.append(values)
        even = row_idx % 2 == 0
        for c in range(1, n_cols + 1):
            _xlsx_data_style(ws.cell(row=row_idx, column=c), even=even)
        ws.row_dimensions[row_idx].height = 18

    _xlsx_thin_border(ws)
    _xlsx_auto_widths(ws)


def _pdf_excel_text_layout(ws, pdf_path: str):
    """Extract text from a non-tabular PDF and lay it out cleanly."""
    import fitz  # PyMuPDF
    from openpyxl.styles import Font, PatternFill, Alignment

    # Column headers
    ws.cell(row=1, column=1, value="Page")
    ws.cell(row=1, column=2, value="Line")
    ws.cell(row=1, column=3, value="Content")
    for c in range(1, 4):
        _xlsx_header_style(ws.cell(row=1, column=c))
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 90

    doc = fitz.open(pdf_path)
    row_idx = 2
    for page_num, page in enumerate(doc, start=1):
        lines = [l.strip() for l in page.get_text().split("\n") if l.strip()]
        for line_num, line in enumerate(lines, start=1):
            even = row_idx % 2 == 0
            pg_cell = ws.cell(row=row_idx, column=1, value=page_num)
            ln_cell = ws.cell(row=row_idx, column=2, value=line_num)
            ct_cell = ws.cell(row=row_idx, column=3, value=line)
            for cell in (pg_cell, ln_cell):
                cell.font = Font(name="Calibri", size=10, color="64748B")
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if even:
                    cell.fill = PatternFill("solid", fgColor="EFF4FB")
            _xlsx_data_style(ct_cell, even=even)
            ws.row_dimensions[row_idx].height = 18
            row_idx += 1
    doc.close()
    _xlsx_thin_border(ws)
